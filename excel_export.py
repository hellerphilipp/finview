# FinView — personal finance manager
# Copyright (C) 2026 Philipp Heller
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

import datetime
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select, text

import db
import queries
from models.finance import Account, Transaction


EXPORT_PATH = Path(tempfile.gettempdir()) / "finview_transactions.xlsx"
TEMPLATE_PATH = Path(__file__).parent / "finview_transactions_template.xlsx"


def export_to_excel(session_factory):
    """Export all data to an Excel workbook and open it."""
    wb = load_workbook(str(TEMPLATE_PATH))

    with session_factory() as session:
        category_paths = queries.get_all_category_paths(session)
        accounts = session.execute(select(Account).order_by(Account.id)).scalars().all()
        transactions = (
            session.execute(select(Transaction).order_by(Transaction.date.desc()))
            .scalars()
            .all()
        )
        category_path_map = queries._build_category_path_map(session)

        _build_categories_sheet(wb, category_paths)
        _build_accounts_sheet(wb, accounts)
        _build_transactions_sheet(wb, transactions, category_path_map, len(category_paths))

    wb.active = wb["Transactions"]
    wb.save(str(EXPORT_PATH))
    subprocess.Popen(["open", str(EXPORT_PATH)])
    return EXPORT_PATH


def _build_categories_sheet(wb, category_paths):
    ws = wb["Categories"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for cat_id, path in category_paths:
        ws.append([cat_id, path])

    ws.tables["CategoriesTable"].ref = f"A1:B{len(category_paths) + 1}"


def _build_accounts_sheet(wb, accounts):
    ws = wb["Accounts"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for acc in accounts:
        ws.append([acc.id, acc.name, acc.currency.value, acc.mapping_spec or ""])

    ws.tables["AccountsTable"].ref = f"A1:D{len(accounts) + 1}"


# Transactions sheet column positions (1-based); col A is a narrow spacer.
_HEADER_ROW = 5
_DATA_START_ROW = 6
_COL_ID = 2
_COL_ACCOUNT_ID = 3
_COL_ACCOUNT = 4
_COL_DESCRIPTION = 5
_COL_ORIGINAL_VALUE = 6
_COL_ORIGINAL_CURRENCY = 7
_COL_VALUE_IN_ACCOUNT_CURRENCY = 8
_COL_DATE = 9
_COL_CATEGORY = 10
_COL_CATEGORY_ID = 11
_COL_REVIEWED = 12
_COL_REVIEWED_AT = 13
_COL_SPLIT_PARENT = 14
_COL_MERGE_PARENT = 15
_NUM_COLS = _COL_MERGE_PARENT  # last column index


def _build_transactions_sheet(wb, transactions, category_path_map, num_categories):
    ws = wb["Transactions"]

    # Clear fake template data (keep title rows 1–5)
    rows_to_delete = ws.max_row - _HEADER_ROW
    if rows_to_delete > 0:
        ws.delete_rows(_DATA_START_ROW, rows_to_delete)

    unlocked = Protection(locked=False)

    for row_idx, tx in enumerate(transactions, start=_DATA_START_ROW):
        account_id_cell = f"{get_column_letter(_COL_ACCOUNT_ID)}{row_idx}"
        category_cell = f"{get_column_letter(_COL_CATEGORY)}{row_idx}"

        account_formula = (
            f'=_xlfn.XLOOKUP({account_id_cell},Accounts!A:A,Accounts!B:B,"")'
        )
        category_id_formula = (
            f'=_xlfn.XLOOKUP({category_cell},Categories!B:B,Categories!A:A,"")'
        )

        category_text = category_path_map.get(tx.category_id, "") if tx.category_id else ""
        reviewed_value = True if tx.reviewed_at else False

        values = [
            (_COL_ID, tx.id),
            (_COL_ACCOUNT_ID, tx.account_id),
            (_COL_ACCOUNT, account_formula),
            (_COL_DESCRIPTION, tx.description),
            (_COL_ORIGINAL_VALUE, float(tx.original_value)),
            (_COL_ORIGINAL_CURRENCY, tx.original_currency.value),
            (_COL_VALUE_IN_ACCOUNT_CURRENCY, float(tx.value_in_account_currency)),
            (_COL_DATE, tx.date),
            (_COL_CATEGORY, category_text),
            (_COL_CATEGORY_ID, category_id_formula),
            (_COL_REVIEWED, reviewed_value),
            (_COL_REVIEWED_AT, tx.reviewed_at),
            (_COL_SPLIT_PARENT, tx.split_parent_id),
            (_COL_MERGE_PARENT, tx.merge_parent_id),
        ]
        for col, val in values:
            ws.cell(row=row_idx, column=col).value = val

        # Unlock editable cells for this new row
        ws.cell(row=row_idx, column=_COL_CATEGORY).protection = unlocked
        ws.cell(row=row_idx, column=_COL_REVIEWED).protection = unlocked

    last_row = _DATA_START_ROW + len(transactions) - 1 if transactions else _HEADER_ROW

    # Update table ref
    ws.tables["TransactionsTable"].ref = (
        f"{get_column_letter(_COL_ID)}{_HEADER_ROW}:"
        f"{get_column_letter(_NUM_COLS)}{last_row}"
    )

    # Extend conditional formatting range if needed.
    # Snapshot items first — mutating a key's sqref changes its hash, which
    # breaks further dict lookups if we iterate and mutate in-place.
    cf_items = list(ws.conditional_formatting._cf_rules.items())
    new_cf_rules = {}
    for cf_obj, rules in cf_items:
        if str(cf_obj.sqref).startswith("F"):
            cf_obj.sqref = f"F{_DATA_START_ROW}:F{last_row}"
        new_cf_rules[cf_obj] = rules
    ws.conditional_formatting._cf_rules = new_cf_rules

    if not transactions:
        return

    # Re-add data validations (openpyxl drops extension-based validations on load)
    data_range = f"{_DATA_START_ROW}:{last_row}"

    if num_categories > 0:
        cat_dv = DataValidation(
            type="list",
            formula1=f"Categories!$B$2:$B${num_categories + 1}",
            allow_blank=True,
        )
        cat_dv.error = "Please select a valid category"
        cat_dv.errorTitle = "Invalid Category"
        ws.add_data_validation(cat_dv)
        cat_dv.add(
            f"{get_column_letter(_COL_CATEGORY)}{_DATA_START_ROW}:"
            f"{get_column_letter(_COL_CATEGORY)}{last_row}"
        )

    rev_dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
    ws.add_data_validation(rev_dv)
    rev_dv.add(
        f"{get_column_letter(_COL_REVIEWED)}{_DATA_START_ROW}:"
        f"{get_column_letter(_COL_REVIEWED)}{last_row}"
    )


def _parse_reviewed(raw) -> bool:
    if isinstance(raw, bool):
        return raw
    if isinstance(raw, str):
        return raw.strip().upper() == "TRUE"
    return bool(raw)


def _resolve_category_id(cat_text, reverse_map: dict) -> int | None:
    if not cat_text or not str(cat_text).strip():
        return None
    return reverse_map.get(str(cat_text).strip())


def import_from_excel(session_factory) -> tuple[int, int, int]:
    """Read back a user-edited Excel export and commit changes to the database.

    Returns (categories_updated, marked_reviewed, marked_unreviewed).
    """
    if not EXPORT_PATH.exists():
        raise FileNotFoundError(f"Export file not found: {EXPORT_PATH}")

    wb = load_workbook(str(EXPORT_PATH), data_only=True)
    ws = wb["Transactions"]

    if "TransactionsTable" not in ws.tables:
        return 0, 0, 0

    table = ws.tables["TransactionsTable"]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)

    header_row = next(
        ws.iter_rows(
            min_row=min_row,
            max_row=min_row,
            min_col=min_col,
            max_col=max_col,
            values_only=True,
        )
    )
    col_index = {name: i for i, name in enumerate(header_row)}

    required = {"id", "category", "reviewed"}
    missing = required - col_index.keys()
    if missing:
        raise ValueError(f"TransactionsTable is missing required columns: {missing}")

    idx_id = col_index["id"]
    idx_category = col_index["category"]
    idx_reviewed = col_index["reviewed"]

    rows_data = []
    for row in ws.iter_rows(
        min_row=min_row + 1,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        tx_id = row[idx_id]
        if tx_id is None:
            continue
        rows_data.append((int(tx_id), row[idx_category], row[idx_reviewed]))

    with session_factory() as session:
        pairs = queries.get_all_category_paths(session)
    reverse_map = {path: cat_id for cat_id, path in pairs}

    now = datetime.datetime.now()
    cat_updates = []
    mark_reviewed = []
    mark_unreviewed = []

    for tx_id, cat_text, raw_reviewed in rows_data:
        cat_updates.append(
            {"id": tx_id, "new_cat": _resolve_category_id(cat_text, reverse_map)}
        )
        if _parse_reviewed(raw_reviewed):
            mark_reviewed.append({"id": tx_id, "now": now})
        else:
            mark_unreviewed.append({"id": tx_id})

    with session_factory() as session:
        cat_result = session.execute(
            text(
                "UPDATE transactions SET category_id = :new_cat "
                "WHERE id = :id AND category_id IS NOT :new_cat"
            ),
            cat_updates,
        )
        categories_updated = cat_result.rowcount

        if mark_reviewed:
            rev_result = session.execute(
                text(
                    "UPDATE transactions SET reviewed_at = :now "
                    "WHERE id = :id AND reviewed_at IS NULL"
                ),
                mark_reviewed,
            )
            marked_reviewed = rev_result.rowcount
        else:
            marked_reviewed = 0

        if mark_unreviewed:
            unrev_result = session.execute(
                text(
                    "UPDATE transactions SET reviewed_at = NULL "
                    "WHERE id = :id AND reviewed_at IS NOT NULL"
                ),
                mark_unreviewed,
            )
            marked_unreviewed = unrev_result.rowcount
        else:
            marked_unreviewed = 0

        session.commit()
        db.mark_dirty()

    return categories_updated, marked_reviewed, marked_unreviewed
